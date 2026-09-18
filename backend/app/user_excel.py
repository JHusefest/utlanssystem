"""Lesing av Excel-ark for import av brukere.

Forventet format (første ark, første rad er overskrifter):

    Elev | Klasse | Brukernavn | Passord

Administratorer importeres aldri – rader som ser ut som administratorer
hoppes over og rapporteres som merknader.
"""

from openpyxl import load_workbook

HEADER_ALIASES: dict[str, str] = {
    "elev": "navn",
    "navn": "navn",
    "fulltnavn": "navn",
    "name": "navn",
    "fullname": "navn",
    "klasse": "klasse",
    "class": "klasse",
    "brukernavn": "brukernavn",
    "username": "brukernavn",
    "bruker": "brukernavn",
    "passord": "passord",
    "password": "passord",
    "rolle": "rolle",
    "role": "rolle",
}

REQUIRED = ("navn", "brukernavn", "passord")
LABELS = {"navn": "Elev", "brukernavn": "Brukernavn", "passord": "Passord"}

ADMIN_WORDS = ("admin", "administrator", "administrer")


def _norm(value) -> str:
    return "".join(ch for ch in str(value or "").lower() if ch.isalnum())


def _text(value) -> str | None:
    if value is None:
        return None
    # Tall i Excel (f.eks. passord «123456») kommer inn som float
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    if not text or text.lower() == "none":
        return None
    return text


def looks_like_admin(name: str | None, username: str | None, role: str | None) -> bool:
    if role and _norm(role) in ("admin", "administrator"):
        return True
    for value in (name, username):
        if value and _norm(value).startswith("admin"):
            return True
    return False


def parse_users_workbook(stream) -> tuple[list[tuple[int, dict]], list[dict], int]:
    """Returnerer (rader, merknader, antall administratorrader som ble hoppet over).

    Hver rad er (radnummer, {"full_name", "school_class", "username", "password"}).
    """
    try:
        wb = load_workbook(stream, data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"Klarte ikke å lese filen: {exc}") from exc

    ws = wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        wb.close()
        return [], [{"row": 1, "message": "Arket er tomt."}], 0

    mapping: dict[int, str] = {}
    for idx, raw in enumerate(header_row):
        key = HEADER_ALIASES.get(_norm(raw))
        if key and key not in mapping.values():
            mapping[idx] = key

    missing = [LABELS[k] for k in REQUIRED if k not in mapping.values()]
    if missing:
        wb.close()
        raise ValueError(
            "Fant ikke kolonnene: "
            + ", ".join(f"«{m}»" for m in missing)
            + ". Arket må ha overskriftene Elev, Klasse, Brukernavn og Passord."
        )

    parsed: list[tuple[int, dict]] = []
    errors: list[dict] = []
    seen: set[str] = set()
    admins_skipped = 0

    for row_no, raw_row in enumerate(rows_iter, start=2):
        values = {
            key: _text(raw_row[idx]) if idx < len(raw_row) else None
            for idx, key in mapping.items()
        }
        if all(v is None for v in values.values()):
            continue

        name = values.get("navn")
        username = values.get("brukernavn")
        password = values.get("passord")

        if looks_like_admin(name, username, values.get("rolle")):
            admins_skipped += 1
            errors.append(
                {
                    "row": row_no,
                    "message": f"«{username or name}» er en administrator – ikke importert.",
                }
            )
            continue

        if not username:
            errors.append({"row": row_no, "message": "Mangler brukernavn – raden ble hoppet over."})
            continue
        if len(username) < 2 or len(username) > 64:
            errors.append(
                {"row": row_no, "message": f"Brukernavnet «{username}» må ha 2–64 tegn – hoppet over."}
            )
            continue
        if not name:
            errors.append({"row": row_no, "message": f"«{username}» mangler navn – hoppet over."})
            continue
        if not password or len(password) < 6:
            errors.append(
                {"row": row_no, "message": f"«{username}» mangler passord eller det er kortere enn 6 tegn – hoppet over."}
            )
            continue

        key = username.lower()
        if key in seen:
            errors.append(
                {"row": row_no, "message": f"Brukernavnet «{username}» finnes flere ganger i filen – hoppet over."}
            )
            continue
        seen.add(key)

        parsed.append(
            (
                row_no,
                {
                    "full_name": name[:120],
                    "school_class": (values.get("klasse") or None),
                    "username": username,
                    "password": password,
                },
            )
        )

    wb.close()
    return parsed, errors, admins_skipped
